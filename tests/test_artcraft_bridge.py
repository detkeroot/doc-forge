# Файл: tests/test_artcraft_bridge.py
"""
Юнит-тесты для модуля интеграции ArtCraft Suite с Doc-Forge.
"""

import os
import unittest
import tempfile
from pathlib import Path
from docx import Document

from src.artcraft_bridge import (
    find_craft_binary,
    is_artcraft_available,
    fast_convert_docx_to_pdf,
    render_page_preview,
    inspect_document_structure,
)


class TestArtCraftBridge(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Создаем временный тестовый DOCX документ
        cls.test_dir = tempfile.TemporaryDirectory()
        cls.docx_path = Path(cls.test_dir.name) / "test_sample.docx"
        doc = Document()
        doc.add_heading("Тестовый документ Doc-Forge", level=1)
        doc.add_paragraph("Это тестовый абзац для проверки работы нативного моста ArtCraft.")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Параметр"
        table.cell(0, 1).text = "Значение"
        table.cell(1, 0).text = "Стандарт"
        table.cell(1, 1).text = "ГОСТ 7.32"
        doc.save(str(cls.docx_path))

    @classmethod
    def tearDownClass(cls):
        cls.test_dir.cleanup()

    def test_find_craft_binary(self):
        """Проверяет обнаружение бинарников ArtCraft в системе."""
        wordcraft_bin = find_craft_binary("wordcraft-cli")
        pdfcraft_bin = find_craft_binary("pdfcraft-cli")
        self.assertIsNotNone(wordcraft_bin, "Бинарник wordcraft-cli должен быть обнаружен")
        self.assertIsNotNone(pdfcraft_bin, "Бинарник pdfcraft-cli должен быть обнаружен")
        self.assertTrue(wordcraft_bin.exists(), "Путь к wordcraft-cli должен существовать")
        self.assertTrue(os.access(wordcraft_bin, os.X_OK), "wordcraft-cli должен быть исполняемым")

    def test_is_artcraft_available(self):
        """Проверяет статус доступности комплекса ArtCraft."""
        self.assertTrue(is_artcraft_available())

    def test_fast_convert_docx_to_pdf(self):
        """Проверяет конвертацию DOCX в PDF через WordCraft."""
        pdf_out = Path(self.test_dir.name) / "test_converted.pdf"
        res_pdf = fast_convert_docx_to_pdf(self.docx_path, pdf_out)
        self.assertTrue(res_pdf.exists(), "Сконвертированный PDF должен существовать")
        self.assertGreater(res_pdf.stat().st_size, 1000, "PDF должен содержать данные")

    def test_render_page_preview_docx(self):
        """Проверяет генерацию PNG превью из DOCX через WordCraft."""
        png_out = Path(self.test_dir.name) / "docx_page1.png"
        res_png = render_page_preview(self.docx_path, page=1, out_path=png_out)
        self.assertTrue(res_png.exists(), "PNG превью DOCX страницы должно существовать")
        self.assertGreater(res_png.stat().st_size, 500, "PNG превью должно содержать пиксели")

    def test_render_page_preview_pdf(self):
        """Проверяет генерацию PNG превью из PDF через PdfCraft."""
        pdf_out = Path(self.test_dir.name) / "test_for_preview.pdf"
        fast_convert_docx_to_pdf(self.docx_path, pdf_out)
        png_out = Path(self.test_dir.name) / "pdf_page1.png"
        res_png = render_page_preview(pdf_out, page=1, out_path=png_out)
        self.assertTrue(res_png.exists(), "PNG превью PDF страницы должно существовать")
        self.assertGreater(res_png.stat().st_size, 500, "PNG превью должно содержать пиксели")

    def test_inspect_document_structure(self):
        """Проверяет извлечение структуры DOCX и PDF через ArtCraft."""
        info_docx = inspect_document_structure(self.docx_path)
        self.assertIn("pages", info_docx)
        self.assertIn("paragraphs", info_docx)
        self.assertGreaterEqual(info_docx["pages"], 1)

        pdf_out = Path(self.test_dir.name) / "test_for_inspect.pdf"
        fast_convert_docx_to_pdf(self.docx_path, pdf_out)
        info_pdf = inspect_document_structure(pdf_out)
        self.assertIn("pages", info_pdf)
        self.assertGreaterEqual(info_pdf["pages"], 1)


if __name__ == "__main__":
    unittest.main()
