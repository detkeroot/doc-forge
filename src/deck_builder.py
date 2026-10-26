# Файл: src/deck_builder.py
"""
Движок сборки нативных, 100% редактируемых презентаций PowerPoint (.pptx)
из структурированного Markdown с автоматическим скачиванием и внедрением изображений.

Инварианты:
1. Запрет растрового запекания: все текстовые блоки, карточки, таблицы - нативные объекты PPTX.
2. Автоматическое скачивание изображений из интернета (URL) с кэшированием в .deck_cache/.
3. Формат 16:9 widescreen (13.333 x 7.5 дюймов).
4. Поддержка заметок докладчика (Presenter Notes).
"""

import os
import re
import urllib.request
import urllib.parse
from pathlib import Path
from typing import List, Dict, Any, Optional, Tuple

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.dml.color import RGBColor


# Цветовая палитра: Современный академический стиль (Academic Clean Slate)
COLOR_BG_WHITE = RGBColor(255, 255, 255)
COLOR_HEADER_NAVY = RGBColor(30, 41, 59)       # #1E293B
COLOR_ACCENT_BLUE = RGBColor(2, 132, 199)      # #0284C7
COLOR_TEXT_PRIMARY = RGBColor(51, 65, 85)      # #334155
COLOR_TEXT_MUTED = RGBColor(100, 116, 139)     # #64748B
COLOR_CARD_BG = RGBColor(248, 250, 252)        # #F8FAFC
COLOR_CARD_BORDER = RGBColor(226, 232, 240)    # #E2E8F0
COLOR_BADGE_BG = RGBColor(224, 242, 254)       # #E0F2FE
COLOR_BADGE_TEXT = RGBColor(3, 105, 161)       # #0369A1
COLOR_SUCCESS_BG = RGBColor(240, 253, 244)     # #F0FDF4
COLOR_SUCCESS_TEXT = RGBColor(21, 128, 61)     # #15803D


class SlideData:
    def __init__(self):
        self.title: str = ""
        self.subtitle: str = ""
        self.header_tag: str = ""
        self.footer_tag: str = ""
        self.bullets: List[str] = []
        self.cards: List[Tuple[str, str]] = []  # (Header, Body)
        self.metrics: List[Tuple[str, str]] = []  # (Value, Label)
        self.table_data: Optional[List[List[str]]] = None
        self.image_path: Optional[str] = None
        self.image_alt: str = ""
        self.notes: str = ""
        self.is_title_slide: bool = False
        self.is_final_slide: bool = False


def download_image_if_url(img_source: str, cache_dir: Path, base_dir: Optional[Path] = None) -> Optional[Path]:
    """
    Скачивает изображение по URL или возвращает путь к локальному файлу.
    """
    if not img_source:
        return None

    # Проверяем локальный файл (абсолютный или относительно base_dir)
    local_p = Path(img_source)
    if local_p.is_absolute() and local_p.exists():
        return local_p
    if base_dir and (base_dir / img_source).exists():
        return base_dir / img_source
    if local_p.exists():
        return local_p

    # Проверяем, URL ли это
    if img_source.startswith("http://") or img_source.startswith("https://"):
        cache_dir.mkdir(parents=True, exist_ok=True)
        parsed = urllib.parse.urlparse(img_source)
        file_ext = Path(parsed.path).suffix.lower()
        if not file_ext or file_ext not in [".png", ".jpg", ".jpeg", ".webp"]:
            file_ext = ".jpg"
        
        url_hash = abs(hash(img_source)) % 10000000
        safe_name = f"img_{url_hash}{file_ext}"
        cached_file = cache_dir / safe_name

        if cached_file.exists() and cached_file.stat().st_size > 100:
            return cached_file

        headers = {
            "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
        }
        try:
            req = urllib.request.Request(img_source, headers=headers)
            with urllib.request.urlopen(req, timeout=12) as response:
                content = response.read()
                if len(content) > 500:
                    cached_file.write_bytes(content)
                    return cached_file
        except Exception as e:
            print(f"[DeckBuilder] Предупреждение: Не удалось скачать изображение '{img_source}': {e}")
            return None

    return None


def parse_markdown_to_slides(md_text: str) -> List[SlideData]:
    """
    Парсит структурированный Markdown на отдельные слайды.
    Разделитель слайдов: '---'
    """
    # Удаляем yaml frontmatter если есть
    if md_text.strip().startswith("---"):
        parts = md_text.strip().split("---", 2)
        if len(parts) >= 3:
            md_text = parts[2].strip()

    raw_slides = re.split(r'\n---\s*\n', md_text)
    slides: List[SlideData] = []

    for idx, raw in enumerate(raw_slides):
        raw = raw.strip()
        if not raw:
            continue

        s = SlideData()
        if idx == 0:
            s.is_title_slide = True

        lines = raw.split("\n")
        in_table = False
        table_rows: List[List[str]] = []

        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue

            # Заметки спикера
            if line_str.startswith("> 💡") or line_str.startswith("> **Заметки") or line_str.startswith("> Заметки"):
                note_clean = re.sub(r'^>\s*(💡|\*\*Заметки.*?\*\*:?|Заметки:?)\s*', '', line_str)
                s.notes += note_clean + " "
                continue
            if line_str.startswith("<!-- notes:") and line_str.endswith("-->"):
                s.notes += line_str[11:-3].strip() + " "
                continue

            # Изображения: ![alt](url_or_path)
            img_match = re.search(r'!\[(.*?)\]\((.*?)\)', line_str)
            if img_match:
                s.image_alt = img_match.group(1)
                s.image_path = img_match.group(2)
                continue

            # Таблицы Markdown
            if line_str.startswith("|") and line_str.endswith("|"):
                if re.match(r'^\|[\s\:\-\|]+$', line_str):
                    continue
                cells = [c.strip() for c in line_str.strip("|").split("|")]
                table_rows.append(cells)
                in_table = True
                continue
            elif in_table:
                in_table = False

            # Заголовки слайда
            if line_str.startswith("# "):
                s.title = line_str[2:].strip()
                continue
            elif line_str.startswith("## "):
                s.title = line_str[3:].strip()
                continue
            elif line_str.startswith("### "):
                s.subtitle = line_str[4:].strip()
                continue

            # Карточки: - **Заголовок:** Описание
            card_match = re.match(r'^[-*]\s*\*\*(.*?)\*\*:?\s*(.*)$', line_str)
            if card_match:
                h = card_match.group(1).strip()
                b = card_match.group(2).strip()
                # Проверяем, не метрика ли это (например, 85%)
                if re.match(r'^\d+[%]?$', h) or "балл" in h.lower():
                    s.metrics.append((h, b))
                else:
                    s.cards.append((h, b))
                continue

            # Обычные буллеты: - Текст
            if line_str.startswith("- ") or line_str.startswith("* "):
                bullet_text = line_str[2:].strip()
                s.bullets.append(bullet_text)
                continue

            # Обычный текст
            if not s.title and idx == 0:
                s.title = line_str
            elif not s.subtitle and (s.title and len(s.bullets) == 0 and len(s.cards) == 0):
                s.subtitle = line_str
            else:
                s.bullets.append(line_str)

        if table_rows:
            s.table_data = table_rows

        if "спасибо за внимание" in s.title.lower() or "вопросы" in s.title.lower():
            s.is_final_slide = True

        slides.append(s)

    return slides


def build_presentation(md_text: str, output_path: Path, cache_dir: Optional[Path] = None):
    """
    Создает презентацию PowerPoint (.pptx) из Markdown текста.
    Все элементы нативные и 100% редактируемые.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    effective_cache_dir = cache_dir or (output_path.parent / ".deck_cache")
    slides_data = parse_markdown_to_slides(md_text)

    for idx, s in enumerate(slides_data, 1):
        slide = prs.slides.add_slide(blank_layout)

        if s.notes:
            notes_tf = slide.notes_slide.notes_text_frame
            notes_tf.text = s.notes.strip()

        if s.is_title_slide:
            _render_title_slide(slide, s, prs.slide_width, prs.slide_height)
            continue

        if s.is_final_slide:
            _render_final_slide(slide, s, prs.slide_width, prs.slide_height)
            continue

        _render_content_slide(slide, s, prs.slide_width, prs.slide_height, effective_cache_dir)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(output_path))
    print(f"[DeckBuilder] Презентация успешно сохранена: {output_path} (Слайдов: {len(slides_data)})")


def _render_title_slide(slide, s: SlideData, width, height):
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), width, Inches(0.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_ACCENT_BLUE
    top_bar.line.color.rgb = COLOR_ACCENT_BLUE

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.0), width - Inches(2.0), height - Inches(2.0)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD_BG
    card.line.color.rgb = COLOR_CARD_BORDER
    card.line.width = Pt(1.5)

    tb = slide.shapes.add_textbox(Inches(1.5), Inches(1.5), width - Inches(3.0), height - Inches(3.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p_title = tf.paragraphs[0]
    p_title.text = s.title
    p_title.font.name = "Times New Roman"
    p_title.font.size = Pt(30)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_HEADER_NAVY
    p_title.space_after = Pt(16)
    p_title.alignment = PP_ALIGN.CENTER

    if s.subtitle:
        p_sub = tf.add_paragraph()
        p_sub.text = s.subtitle
        p_sub.font.name = "Times New Roman"
        p_sub.font.size = Pt(18)
        p_sub.font.color.rgb = COLOR_ACCENT_BLUE
        p_sub.font.bold = True
        p_sub.space_after = Pt(24)
        p_sub.alignment = PP_ALIGN.CENTER

    for b in s.bullets:
        p_b = tf.add_paragraph()
        p_b.text = b
        p_b.font.name = "Times New Roman"
        p_b.font.size = Pt(15)
        p_b.font.color.rgb = COLOR_TEXT_PRIMARY
        p_b.space_after = Pt(8)
        p_b.alignment = PP_ALIGN.CENTER


def _render_final_slide(slide, s: SlideData, width, height):
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), width, Inches(0.15))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_ACCENT_BLUE
    top_bar.line.color.rgb = COLOR_ACCENT_BLUE

    card = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.5), width - Inches(3.0), height - Inches(3.0)
    )
    card.fill.solid()
    card.fill.fore_color.rgb = COLOR_CARD_BG
    card.line.color.rgb = COLOR_CARD_BORDER
    card.line.width = Pt(1.5)

    tb = slide.shapes.add_textbox(Inches(2.0), Inches(2.2), width - Inches(4.0), height - Inches(4.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p1 = tf.paragraphs[0]
    p1.text = s.title or "Спасибо за внимание!"
    p1.font.name = "Times New Roman"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_HEADER_NAVY
    p1.alignment = PP_ALIGN.CENTER
    p1.space_after = Pt(20)

    p2 = tf.add_paragraph()
    p2.text = s.subtitle or "Готов ответить на ваши вопросы"
    p2.font.name = "Times New Roman"
    p2.font.size = Pt(22)
    p2.font.color.rgb = COLOR_ACCENT_BLUE
    p2.alignment = PP_ALIGN.CENTER
    p2.space_after = Pt(24)

    for b in s.bullets:
        p_b = tf.add_paragraph()
        p_b.text = b
        p_b.font.name = "Times New Roman"
        p_b.font.size = Pt(16)
        p_b.font.color.rgb = COLOR_TEXT_MUTED
        p_b.alignment = PP_ALIGN.CENTER


def _render_content_slide(slide, s: SlideData, width, height, cache_dir: Path):
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), width, Inches(0.08))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_ACCENT_BLUE
    top_bar.line.color.rgb = COLOR_ACCENT_BLUE

    header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), width - Inches(1.6), Inches(1.1))
    htf = header_box.text_frame
    htf.word_wrap = True
    htf.margin_left = Inches(0)
    htf.margin_top = Inches(0)

    hp = htf.paragraphs[0]
    hp.text = s.title
    hp.font.name = "Times New Roman"
    hp.font.size = Pt(24)
    hp.font.bold = True
    hp.font.color.rgb = COLOR_HEADER_NAVY

    if s.subtitle:
        hsub = htf.add_paragraph()
        hsub.text = s.subtitle
        hsub.font.name = "Times New Roman"
        hsub.font.size = Pt(14)
        hsub.font.color.rgb = COLOR_ACCENT_BLUE
        hsub.space_before = Pt(4)

    div_line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.5), width - Inches(1.6), Inches(0.02))
    div_line.fill.solid()
    div_line.fill.fore_color.rgb = COLOR_CARD_BORDER
    div_line.line.color.rgb = COLOR_CARD_BORDER

    img_file = download_image_if_url(s.image_path, cache_dir, base_dir=cache_dir.parent) if s.image_path else None

    content_top = Inches(1.7)
    content_height = height - Inches(2.2)

    if img_file:
        content_width = Inches(7.6)
        img_left = Inches(8.7)
        img_width = Inches(3.8)
    else:
        content_width = width - Inches(1.6)
        img_left = None
        img_width = None

    if s.table_data and len(s.table_data) > 1:
        _render_native_table(slide, s.table_data, Inches(0.8), content_top, content_width, content_height)
    else:
        _render_cards_and_bullets(slide, s, Inches(0.8), content_top, content_width, content_height)

    if img_file:
        try:
            img_frame = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE, img_left - Inches(0.1), content_top, img_width + Inches(0.2), content_height
            )
            img_frame.fill.solid()
            img_frame.fill.fore_color.rgb = COLOR_CARD_BG
            img_frame.line.color.rgb = COLOR_CARD_BORDER
            img_frame.line.width = Pt(1)

            pic = slide.shapes.add_picture(
                str(img_file), img_left, content_top + Inches(0.15), width=img_width
            )
            if pic.height > (content_height - Inches(0.5)):
                pic.height = content_height - Inches(0.5)
        except Exception as e:
            print(f"[DeckBuilder] Ошибка вставки картинки: {e}")


def _render_cards_and_bullets(slide, s: SlideData, left, top, width, height):
    curr_top = top
    curr_height = height

    if s.metrics:
        badge_count = len(s.metrics)
        col_w = (width - Inches(0.3 * (badge_count - 1))) / badge_count
        badge_h = Inches(1.4)

        for i, (val, lbl) in enumerate(s.metrics):
            bx = left + i * (col_w + Inches(0.3))
            b_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, curr_top, col_w, badge_h)
            b_shape.fill.solid()
            b_shape.fill.fore_color.rgb = COLOR_BADGE_BG
            b_shape.line.color.rgb = COLOR_ACCENT_BLUE
            b_shape.line.width = Pt(1.5)

            btf = b_shape.text_frame
            btf.word_wrap = True
            p_val = btf.paragraphs[0]
            p_val.text = val
            p_val.font.name = "Times New Roman"
            p_val.font.size = Pt(26)
            p_val.font.bold = True
            p_val.font.color.rgb = COLOR_BADGE_TEXT
            p_val.alignment = PP_ALIGN.CENTER

            p_lbl = btf.add_paragraph()
            p_lbl.text = lbl
            p_lbl.font.name = "Times New Roman"
            p_lbl.font.size = Pt(12)
            p_lbl.font.color.rgb = COLOR_TEXT_PRIMARY
            p_lbl.alignment = PP_ALIGN.CENTER
            p_lbl.space_before = Pt(4)

        curr_top += badge_h + Inches(0.25)
        curr_height -= (badge_h + Inches(0.25))

    if s.cards:
        card_count = len(s.cards)
        card_h = (curr_height - Inches(0.2 * (card_count - 1))) / max(card_count, 1)

        for idx, (header, body) in enumerate(s.cards):
            cy = curr_top + idx * (card_h + Inches(0.2))
            c_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, cy, width, card_h)
            c_shape.fill.solid()
            c_shape.fill.fore_color.rgb = COLOR_CARD_BG
            c_shape.line.color.rgb = COLOR_CARD_BORDER
            c_shape.line.width = Pt(1)

            ctf = c_shape.text_frame
            ctf.word_wrap = True
            ctf.margin_left = Inches(0.25)
            ctf.margin_right = Inches(0.25)
            ctf.margin_top = Inches(0.12)
            ctf.margin_bottom = Inches(0.12)

            p_h = ctf.paragraphs[0]
            p_h.text = header
            p_h.font.name = "Times New Roman"
            p_h.font.size = Pt(15)
            p_h.font.bold = True
            p_h.font.color.rgb = COLOR_ACCENT_BLUE

            if body:
                p_b = ctf.add_paragraph()
                p_b.text = body
                p_b.font.name = "Times New Roman"
                p_b.font.size = Pt(13)
                p_b.font.color.rgb = COLOR_TEXT_PRIMARY
                p_b.space_before = Pt(3)

    elif s.bullets:
        tb = slide.shapes.add_textbox(left, curr_top, width, curr_height)
        tf = tb.text_frame
        tf.word_wrap = True

        for idx, b in enumerate(s.bullets):
            p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            p.text = f"•  {b}"
            p.font.name = "Times New Roman"
            p.font.size = Pt(15)
            p.font.color.rgb = COLOR_TEXT_PRIMARY
            p.space_after = Pt(12)


def _render_native_table(slide, table_data: List[List[str]], left, top, width, height):
    rows = len(table_data)
    cols = max(len(r) for r in table_data)

    table_shape = slide.shapes.add_table(rows, cols, left, top, width, min(height, Inches(0.55 * rows)))
    table = table_shape.table

    for r_idx, row in enumerate(table_data):
        for c_idx, cell_text in enumerate(row):
            if c_idx < cols:
                cell = table.cell(r_idx, c_idx)
                cell.text = cell_text
                for p in cell.text_frame.paragraphs:
                    p.font.name = "Times New Roman"
                    p.font.size = Pt(12)
                    if r_idx == 0:
                        p.font.bold = True
                        p.font.color.rgb = COLOR_HEADER_NAVY
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = COLOR_BADGE_BG
                    else:
                        p.font.color.rgb = COLOR_TEXT_PRIMARY
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = COLOR_BG_WHITE if r_idx % 2 == 1 else COLOR_CARD_BG
